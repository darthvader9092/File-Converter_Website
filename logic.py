import os
from pydub import AudioSegment
from PIL import Image
from fpdf import FPDF
from docx import Document
from weasyprint import HTML

def handle_conversion(input_path, category, target_format, output_folder):
    filename, _ = os.path.splitext(os.path.basename(input_path))
    output_path = os.path.join(output_folder, f"{filename}.{target_format}")

    try:
        if category == "audio":
            if not target_format in ["mp3", "wav", "ogg", "aac"]:
                raise ValueError(f"Unsupported audio format: {target_format}")
            audio = AudioSegment.from_file(input_path)
            audio.export(output_path, format=target_format)

        elif category == "image":
            if not target_format in ["png", "jpg", "jpeg", "bmp", "gif"]:
                raise ValueError(f"Unsupported image format: {target_format}")
            image = Image.open(input_path)
            image.save(output_path, format=target_format.upper())

        elif category == "document":
            if target_format == "pdf":
                if input_path.endswith(".txt"):
                    pdf = FPDF()
                    pdf.add_page()
                    pdf.set_font("Arial", size=12)
                    with open(input_path, "r") as file:
                        for line in file:
                            pdf.cell(0, 10, txt=line.strip(), ln=True)
                    pdf.output(output_path)
                elif input_path.endswith(".html"):
                    HTML(filename=input_path).write_pdf(output_path)
                else:
                    raise ValueError("Unsupported document type for PDF conversion.")
            elif target_format == "docx":
                doc = Document()
                if input_path.endswith(".txt"):
                    with open(input_path, "r") as file:
                        doc.add_paragraph(file.read())
                else:
                    raise ValueError("Only .txt files can be converted to docx.")
                doc.save(output_path)
            else:
                raise ValueError(f"Unsupported target document format: {target_format}")

        else:
            raise ValueError("Unsupported category. Choose from 'audio', 'image', or 'document'.")

    except Exception as e:
        print(f"Error during conversion: {e}")
        return None

    return output_path

